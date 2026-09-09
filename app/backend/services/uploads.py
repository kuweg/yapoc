"""Upload store — disk persistence, dedup, indexing and resolution for chat
attachments (Phase 1 of the attachment system).

Layout under ``settings.upload_dir`` (``data/uploads/``):
  {YYYY}/{MM}/{DD}/{uuid.hex}.{safe_ext}   the stored bytes
  uploads.json                              id -> metadata index (atomic write)
  .thumbs/{id}.jpg                          cached 320x320 JPEG thumbnails
  .vision/{id}.txt                          cached VL OCR/caption text (Phase 3)

Every record is owner-scoped; ``resolve_upload`` drops anything the caller does
not own so attachments can't leak across users.
"""
from __future__ import annotations

import hashlib
import io
import json
import mailbox
import mimetypes
import os
import re
import tarfile
import threading
import uuid
import zipfile
from email import policy
from email.parser import BytesParser
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Optional

from app.config import settings

# libmagic is optional — fall back to extension-based detection if unavailable.
try:
    import magic as _libmagic  # python-magic
    _HAVE_MAGIC = True
except Exception:  # pragma: no cover - environment dependent
    _HAVE_MAGIC = False

try:
    from PIL import Image, ImageOps
    _HAVE_PIL = True
except Exception:  # pragma: no cover
    _HAVE_PIL = False

THUMB_SIZE = (320, 320)

# Reject obviously dangerous executables/scripts (defence in depth — these are
# never something a chat attachment legitimately needs).
_DANGEROUS_EXT = {
    ".exe", ".dll", ".bat", ".cmd", ".com", ".scr", ".msi", ".vbs", ".vbe",
    ".js", ".jse", ".ps1", ".psm1", ".jsp", ".asp", ".aspx", ".jar", ".sh",
    ".bash", ".csh", ".pif", ".cpl", ".hta", ".wsf",
}
_DANGEROUS_MIME = {
    "application/x-msdownload", "application/x-dosexec", "application/x-executable",
    "application/x-sh", "application/x-shellscript", "application/javascript",
    "text/javascript", "application/x-msdos-program", "application/vnd.microsoft.portable-executable",
}

_lock = threading.Lock()


# ── paths ────────────────────────────────────────────────────────────────────
def _root() -> Path:
    return settings.upload_dir


def _index_path() -> Path:
    return _root() / "uploads.json"


def _safe_ext(filename: str, mime: str) -> str:
    ext = Path(filename or "").suffix.lower()
    ext = re.sub(r"[^a-z0-9.]", "", ext)[:10]
    if not ext:
        guessed = mimetypes.guess_extension(mime or "") or ""
        ext = re.sub(r"[^a-z0-9.]", "", guessed.lower())[:10]
    return ext or ".bin"


# ── mime / validation ────────────────────────────────────────────────────────
def detect_mime(content: bytes, filename: str) -> str:
    if _HAVE_MAGIC:
        try:
            m = _libmagic.from_buffer(content[:8192], mime=True)
            if m and m != "application/octet-stream":
                return m
        except Exception:
            pass
    guessed, _ = mimetypes.guess_type(filename or "")
    return guessed or "application/octet-stream"


def is_dangerous(mime: str, filename: str) -> bool:
    ext = Path(filename or "").suffix.lower()
    return ext in _DANGEROUS_EXT or (mime or "").lower() in _DANGEROUS_MIME


def image_dimensions(content: bytes, mime: str) -> tuple[Optional[int], Optional[int]]:
    if not (mime or "").startswith("image/") or not _HAVE_PIL:
        return None, None
    try:
        im = Image.open(io.BytesIO(content))
        im = ImageOps.exif_transpose(im)  # honour EXIF orientation
        return int(im.width), int(im.height)
    except Exception:
        return None, None


# ── index (atomic) ───────────────────────────────────────────────────────────
def _load_index() -> dict[str, dict[str, Any]]:
    p = _index_path()
    if not p.exists():
        return {}
    try:
        return json.loads(p.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError):
        bak = p.with_suffix(".json.bak")
        if bak.exists():
            try:
                return json.loads(bak.read_text(encoding="utf-8"))
            except Exception:
                pass
        return {}


def _save_index(index: dict[str, dict[str, Any]]) -> None:
    root = _root()
    root.mkdir(parents=True, exist_ok=True)
    p = _index_path()
    if p.exists():
        try:
            p.replace(p.with_suffix(".json.bak"))
        except OSError:
            pass
    tmp = p.with_suffix(".json.tmp")
    tmp.write_text(json.dumps(index, indent=2), encoding="utf-8")
    os.replace(tmp, p)


# ── public meta projection (what the API returns) ────────────────────────────
def public_meta(rec: dict[str, Any]) -> dict[str, Any]:
    out = {
        "id": rec["id"],
        "name": rec.get("name") or rec.get("original_name"),
        "mime": rec.get("mime"),
        "size": rec.get("size"),
        "hash": rec.get("hash"),
        "is_duplicate": rec.get("is_duplicate", False),
    }
    if rec.get("width") is not None:
        out["width"] = rec["width"]
    if rec.get("height") is not None:
        out["height"] = rec["height"]
    return out


# ── store / resolve ──────────────────────────────────────────────────────────
def store_upload(filename: str, content: bytes, owner: str, client_ip: str = "") -> dict[str, Any]:
    """Validate + persist one file, dedup per owner, index it. Returns the record
    (with ``is_duplicate`` reflecting whether the owner already had these bytes).

    Raises ValueError on validation failure (caller maps to HTTP 4xx).
    """
    if len(content) == 0:
        raise ValueError("Empty file")

    mime = detect_mime(content, filename)
    if is_dangerous(mime, filename):
        raise ValueError(f"Disallowed file type: {filename} ({mime})")

    digest = hashlib.sha256(content).hexdigest()
    now = datetime.now(timezone.utc).isoformat()
    width, height = image_dimensions(content, mime)

    with _lock:
        index = _load_index()
        # Dedup: same owner + same bytes → return the existing record.
        for rec in index.values():
            if rec.get("owner") == owner and rec.get("hash") == digest:
                rec = dict(rec)
                rec["is_duplicate"] = True
                rec["last_accessed"] = now
                index[rec["id"]] = {**index[rec["id"]], "last_accessed": now}
                _save_index(index)
                return rec

        file_id = uuid.uuid4().hex
        ext = _safe_ext(filename, mime)
        dt = datetime.now(timezone.utc)
        rel_dir = Path(f"{dt:%Y}") / f"{dt:%m}" / f"{dt:%d}"
        abs_dir = _root() / rel_dir
        abs_dir.mkdir(parents=True, exist_ok=True)
        stored = abs_dir / f"{file_id}{ext}"
        stored.write_bytes(content)

        rec = {
            "id": file_id,
            "name": Path(filename or "file").name or "file",
            "original_name": filename,
            "mime": mime,
            "size": len(content),
            "hash": digest,
            "width": width,
            "height": height,
            "path": str(stored.relative_to(_root())),
            "owner": owner,
            "client_ip": client_ip,
            "uploaded_at": now,
            "last_accessed": now,
            "is_duplicate": False,
        }
        index[file_id] = rec
        _save_index(index)
        return rec


def resolve_upload(file_id: str, owner: Optional[str] = None) -> Optional[dict[str, Any]]:
    """Look up a record by id, owner-scoped. Returns None if missing or not owned
    (admin/None owner bypasses the scope check)."""
    with _lock:
        index = _load_index()
        rec = index.get(file_id)
        if not rec:
            return None
        if owner is not None and rec.get("owner") not in (owner, None, ""):
            return None
        return dict(rec)


def resolve_upload_by_ref(ref: str, owner: str) -> Optional[dict[str, Any]]:
    """Resolve a raw upload id or an ``@file:<id|name>`` reference for an owner."""
    value = (ref or "").strip()
    if value.startswith("@file:"):
        value = value[len("@file:"):].strip()
    record = resolve_upload(value, owner=owner)
    if record:
        return record
    if not value:
        return None
    with _lock:
        for candidate in _load_index().values():
            if candidate.get("owner") == owner and candidate.get("name") == value:
                return dict(candidate)
    return None


def upload_path(rec: dict[str, Any]) -> Path:
    return _root() / rec["path"]


# ── thumbnails ───────────────────────────────────────────────────────────────
def get_thumbnail(rec: dict[str, Any]) -> Optional[Path]:
    """Return a cached 320x320 JPEG thumbnail for an image record, generating it
    on first request. None for non-images or when Pillow is unavailable."""
    if not (rec.get("mime") or "").startswith("image/") or not _HAVE_PIL:
        return None
    thumbs = _root() / ".thumbs"
    thumbs.mkdir(parents=True, exist_ok=True)
    out = thumbs / f"{rec['id']}.jpg"
    if out.exists():
        return out
    src = upload_path(rec)
    if not src.exists():
        return None
    try:
        im = Image.open(src)
        im = ImageOps.exif_transpose(im).convert("RGB")
        im.thumbnail(THUMB_SIZE)
        im.save(out, "JPEG", quality=82)
        return out
    except Exception:
        return None


def touch_accessed(file_id: str) -> None:
    with _lock:
        index = _load_index()
        rec = index.get(file_id)
        if rec:
            rec["last_accessed"] = datetime.now(timezone.utc).isoformat()
            _save_index(index)


def delete_upload(file_id: str, owner: Optional[str] = None) -> bool:
    """Remove an upload record and its stored file, thumbnail and vision cache.

    Owner-scoped (mirrors ``resolve_upload``); index removal is the source of
    truth and on-disk unlink failures are ignored. Returns True if removed.
    """
    with _lock:
        index = _load_index()
        rec = index.get(file_id)
        if not rec:
            return False
        if owner is not None and rec.get("owner") not in (owner, None, ""):
            return False
        del index[file_id]
        _save_index(index)

    stored = upload_path(rec)
    if stored.exists():
        try:
            stored.unlink()
        except OSError:
            pass

    thumb = _root() / ".thumbs" / f"{file_id}.jpg"
    if thumb.exists():
        try:
            thumb.unlink()
        except OSError:
            pass

    vision = _root() / ".vision" / f"{file_id}.txt"
    if vision.exists():
        try:
            vision.unlink()
        except OSError:
            pass

    return True


def rename_upload(file_id: str, new_name: str, owner: Optional[str] = None) -> Optional[dict[str, Any]]:
    """Rename an upload's display name (owner-scoped). Returns the updated record
    or None if the file is missing or not owned. The on-disk file keeps its
    stored path; only the display ``name`` (and ``original_name``) change."""
    new_name = (new_name or "").strip()
    if not new_name:
        return None
    with _lock:
        index = _load_index()
        rec = index.get(file_id)
        if not rec:
            return None
        if owner is not None and rec.get("owner") not in (owner, None, ""):
            return None
        rec["name"] = new_name
        rec["original_name"] = new_name
        _save_index(index)
        return dict(rec)


# ── message injection (Phase 1 — image_read marker path + inline text) ───────
TEXT_BUDGET = 24_000  # total inlined chars across all attachments in one message
PDF_MAX_PAGES = 100  # bound CPU/work for malformed or unusually large PDFs
PDF_TEXT_BUDGET = 24_000  # bound extraction before the message-level budget is applied
XLSX_MAX_ROWS = 500  # cap rows read per sheet
XLSX_MAX_CELLS = 20_000  # hard cap on total cells across all sheets
ARCHIVE_MAX_ENTRIES = 500  # cap entries listed in an archive
ARCHIVE_MAX_TOTAL_BYTES = 50 * 1024 * 1024  # 50MB cap on uncompressed archive size
EMAIL_MAX_CHARS = 24_000  # per-message text budget
ICS_MAX_EVENTS = 200
_TEXT_MIMES = ("text/", "application/json", "application/xml", "application/x-yaml")


def project_rel_path(rec: dict[str, Any]) -> str:
    """Path of the stored file relative to project_root, so the agent's
    ``image_read`` / ``file_read`` tools (which sandbox to project_root) can load
    it directly."""
    return f"data/uploads/{rec['path']}"


def _extract_xlsx(path: Path) -> Optional[str]:
    """Read a bounded, value-only workbook preview as Markdown tables."""
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        if not workbook.worksheets:
            return None
        total_cells = 0
        sections: list[str] = []
        exhausted = False
        for sheet in workbook.worksheets:
            rows: list[list[str]] = []
            for row_index, row in enumerate(sheet.iter_rows(values_only=True)):
                if row_index >= XLSX_MAX_ROWS or exhausted:
                    break
                total_cells += len(row)
                if total_cells > XLSX_MAX_CELLS:
                    exhausted = True
                    break
                values = ["" if value is None else str(value) for value in row]
                if any(value.strip() for value in values):
                    rows.append(values)
            if not rows:
                continue
            width = max(len(row) for row in rows)
            normalized = [row + [""] * (width - len(row)) for row in rows]
            header = normalized[0]
            table = [
                f"## Sheet: {sheet.title}",
                "| " + " | ".join(value.replace("|", "\\|") for value in header) + " |",
                "| " + " | ".join("---" for _ in header) + " |",
            ]
            table.extend(
                "| " + " | ".join(value.replace("|", "\\|") for value in row) + " |"
                for row in normalized[1:]
            )
            sections.append("\n".join(table))
            if exhausted:
                break
        workbook.close()
        return "\n\n".join(sections) or None
    except Exception:
        return None


def _walk_pptx_shapes(shapes: Any) -> list[str]:
    """Collect visible text from shapes, tables, and nested group shapes."""
    text: list[str] = []
    for shape in shapes:
        if getattr(shape, "has_text_frame", False):
            text.extend(paragraph.text for paragraph in shape.text_frame.paragraphs if paragraph.text)
        if getattr(shape, "has_table", False):
            for row in shape.table.rows:
                values = [cell.text.strip() for cell in row.cells]
                if any(values):
                    text.append(" | ".join(values))
        if hasattr(shape, "shapes"):
            text.extend(_walk_pptx_shapes(shape.shapes))
    return text


def _extract_pptx(path: Path) -> Optional[str]:
    try:
        from pptx import Presentation

        presentation = Presentation(str(path))
        sections: list[str] = []
        remaining = 24_000
        for index, slide in enumerate(presentation.slides, start=1):
            if remaining <= 0:
                break
            body = "\n".join(_walk_pptx_shapes(slide.shapes))
            section = f"## Slide {index}\n{body}" if body else f"## Slide {index}"
            sections.append(section[:remaining])
            remaining -= len(section)
        return "\n\n".join(sections) or None
    except Exception:
        return None


def _email_message_text(message: Any) -> str:
    headers = "\n".join(
        f"{header}: {message.get(header, '')}" for header in ("From", "To", "Subject", "Date")
    )
    plain: list[str] = []
    html: list[str] = []
    for part in message.walk() if message.is_multipart() else [message]:
        if part.is_multipart() or part.get_content_disposition() == "attachment":
            continue
        content_type = part.get_content_type()
        if content_type not in ("text/plain", "text/html"):
            continue
        try:
            payload = part.get_content()
            if not isinstance(payload, str):
                payload = str(payload)
        except Exception:
            continue
        if content_type == "text/plain":
            plain.append(payload)
        else:
            html.append(re.sub(r"<[^>]+>", " ", payload))
    body = "\n".join(plain) or "\n".join(html)
    return f"{headers}\n\n{body}".strip()


def _extract_email(path: Path, mime: str, name: str) -> Optional[str]:
    try:
        if name.lower().endswith(".mbox"):
            messages: list[str] = []
            remaining = EMAIL_MAX_CHARS
            box = mailbox.mbox(str(path), create=False)
            try:
                for index, message in enumerate(box, start=1):
                    if index > 50 or remaining <= 0:
                        break
                    block = f"--- Message {index} ---\n{_email_message_text(message)}"
                    messages.append(block[:remaining])
                    remaining -= len(block)
            finally:
                box.close()
            return "\n\n".join(messages) or None
        message = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
        return _email_message_text(message)[:EMAIL_MAX_CHARS] or None
    except Exception:
        return None


def _extract_ics(path: Path) -> Optional[str]:
    try:
        from icalendar import Calendar

        calendar = Calendar.from_ical(path.read_bytes())
        events: list[str] = []
        for component in calendar.walk("VEVENT"):
            if len(events) >= ICS_MAX_EVENTS:
                break
            fields: list[str] = []
            for field in ("SUMMARY", "DTSTART", "DTEND", "LOCATION", "DESCRIPTION"):
                value = component.get(field)
                if value is not None:
                    fields.append(f"{field}: {value}")
            events.append("\n".join(fields) or "VEVENT")
        return "\n\n".join(events) or None
    except Exception:
        return None


def _unsafe_archive_path(member_name: str) -> bool:
    normalized = member_name.replace("\\", "/")
    return normalized.startswith("/") or any(part == ".." for part in normalized.split("/"))


def _extract_archive_listing(path: Path, mime: str, name: str) -> Optional[str]:
    """List archive metadata only; never extract or write archive members."""
    try:
        lower_name = name.lower()
        entries: list[tuple[str, int, bool, bool]] = []
        if lower_name.endswith(".zip"):
            with zipfile.ZipFile(path) as archive:
                for info in archive.infolist()[:ARCHIVE_MAX_ENTRIES]:
                    entries.append((info.filename, info.file_size, info.is_dir(), _unsafe_archive_path(info.filename)))
        elif lower_name.endswith((".tar", ".tar.gz", ".tgz")):
            with tarfile.open(path, mode="r:*") as archive:
                for member in archive.getmembers()[:ARCHIVE_MAX_ENTRIES]:
                    entries.append((member.name, member.size, member.isdir(), _unsafe_archive_path(member.name)))
        elif lower_name.endswith(".7z"):
            import py7zr

            with py7zr.SevenZipFile(path, mode="r") as archive:
                for entry in archive.list()[:ARCHIVE_MAX_ENTRIES]:
                    entry_name = getattr(entry, "filename", "")
                    entries.append((entry_name, int(getattr(entry, "uncompressed", 0) or 0), bool(getattr(entry, "is_directory", False)), _unsafe_archive_path(entry_name)))
        else:
            return None
        total_bytes = sum(size for _, size, _, _ in entries)
        lines = [f"Archive: {name} ({len(entries)} entries, total {total_bytes} bytes)"]
        for member_name, size, is_dir, unsafe in entries:
            suffix = " [directory]" if is_dir else ""
            if unsafe:
                suffix += " skipped (unsafe path)"
            lines.append(f"  {member_name}  ({size} bytes){suffix}")
        if total_bytes > ARCHIVE_MAX_TOTAL_BYTES:
            lines.append("Extraction refused: total uncompressed size exceeds 50MB safety cap.")
        lines.append("Listing only — no contents extracted. Use the archive_inspect tool to extract a specific safe entry.")
        return "\n".join(lines)
    except Exception:
        return None


def _extract_text(rec: dict[str, Any]) -> Optional[str]:
    """Best-effort text extraction for text/code/markdown/docx/PDF attachments.
    Returns None when extraction is unsupported, unsafe, or fails."""
    mime = (rec.get("mime") or "").lower()
    name = (rec.get("name") or "").lower()
    path = upload_path(rec)
    if not path.exists():
        return None
    is_text = mime.startswith(_TEXT_MIMES) or name.endswith(
        (".txt", ".md", ".markdown", ".json", ".yaml", ".yml", ".csv", ".log",
         ".py", ".js", ".ts", ".tsx", ".jsx", ".html", ".css", ".sh", ".toml", ".ini")
    )
    if is_text:
        try:
            return path.read_text(encoding="utf-8", errors="replace")
        except OSError:
            return None
    if name.endswith(".docx"):
        try:
            import io as _io
            import docx  # python-docx
            doc = docx.Document(_io.BytesIO(path.read_bytes()))
            return "\n".join(p.text for p in doc.paragraphs)
        except Exception:
            return None
    if mime == "application/pdf" or name.endswith(".pdf"):
        try:
            from pypdf import PdfReader

            # pypdf reads directly from disk; do not execute PDF JavaScript or
            # render embedded content. Extraction is capped to keep uploads
            # from consuming unbounded CPU or context budget.
            reader = PdfReader(str(path), strict=False)
            if reader.is_encrypted:
                try:
                    if reader.decrypt("") == 0:
                        return None
                except Exception:
                    return None
            pages: list[str] = []
            remaining = PDF_TEXT_BUDGET
            for page in reader.pages[:PDF_MAX_PAGES]:
                if remaining <= 0:
                    break
                try:
                    text = page.extract_text() or ""
                except Exception:
                    continue
                if text:
                    pages.append(text[:remaining])
                    remaining -= len(text)
            return "\n\n".join(pages)
        except Exception:
            return None
    if name.endswith(".xlsx") or mime == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        try:
            return _extract_xlsx(path)
        except Exception:
            return None
    if name.endswith(".pptx") or mime == "application/vnd.openxmlformats-officedocument.presentationml.presentation":
        try:
            return _extract_pptx(path)
        except Exception:
            return None
    if name.endswith(".eml") or mime == "message/rfc822" or name.endswith(".mbox"):
        try:
            return _extract_email(path, mime, name)
        except Exception:
            return None
    if name.endswith(".ics") or mime == "text/calendar":
        try:
            return _extract_ics(path)
        except Exception:
            return None
    if name.endswith((".zip", ".tar", ".tar.gz", ".tgz", ".7z")):
        try:
            return _extract_archive_listing(path, mime, name)
        except Exception:
            return None
    return None


def build_attachment_injection(ids: list[str], owner: str) -> tuple[str, list[dict[str, Any]]]:
    """Resolve attachment IDs (owner-scoped) and build the text to append to the
    user message: an ``image_read`` marker per image, inline text for
    text/docx, and a note for formats not yet extracted. Returns
    ``(suffix, meta_list)`` where meta_list is the public metadata for the SSE
    ``attachments`` event + bubble rendering.
    """
    parts: list[str] = []
    meta: list[dict[str, Any]] = []
    budget = TEXT_BUDGET
    for fid in ids or []:
        rec = resolve_upload(fid, owner=owner)
        if not rec:
            continue
        meta.append(public_meta(rec))
        mime = (rec.get("mime") or "")
        name = rec.get("name") or "file"
        rel = project_rel_path(rec)
        if mime.startswith("image/"):
            # Master loads it via image_read; normalize.py formats per-adapter.
            parts.append(f"[📎 photo attached: {rel}]")
            continue
        text = _extract_text(rec)
        if text is not None:
            if budget <= 0:
                parts.append(f"\n\n[Attachment omitted (budget exceeded): {name}]")
                continue
            snippet = text[:budget]
            budget -= len(snippet)
            trunc = "\n[… truncated]" if len(text) > len(snippet) else ""
            ext = name.rsplit(".", 1)[-1] if "." in name else ""
            fence = f"```{ext}\n" if ext and ext not in ("txt", "md", "markdown") else ""
            close = "```" if fence else ""
            parts.append(f"\n\n--- Attachment: {name} ---\n{fence}{snippet}{trunc}\n{close}\n--- end {name} ---")
        else:
            parts.append(f"\n\n[Attachment: {name} ({mime}) — content extraction not available yet]")
    return ("\n".join(parts) if parts else ""), meta
